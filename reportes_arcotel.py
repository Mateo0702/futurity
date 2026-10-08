"""
Módulo de generación de reportes regulatorios mensuales para ARCOTEL.
Genera los 3 reportes oficiales con sus formatos institucionales:
1. GPON: Tiempo Promedio de Reparación de Averías (Técnicas) - 11 columnas
2. HFC: Formulario SAV-Q-001 (Requerimientos de Atención de los Suscriptores) - 8 columnas
3. VELOCIDAD: Porcentaje de Reclamos por Capacidad del Canal de Acceso - 10 columnas
"""

import io
import re
import zipfile
from datetime import datetime, date
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import db_config

MESES_NOMBRES = {
    1: 'ENERO', 2: 'FEBRERO', 3: 'MARZO', 4: 'ABRIL',
    5: 'MAYO', 6: 'JUNIO', 7: 'JULIO', 8: 'AGOSTO',
    9: 'SEPTIEMBRE', 10: 'OCTUBRE', 11: 'NOVIEMBRE', 12: 'DICIEMBRE'
}

MESES_ABREV = {
    1: 'ENE', 2: 'FEB', 3: 'MAR', 4: 'ABR',
    5: 'MAY', 6: 'JUN', 7: 'JUL', 8: 'AGO',
    9: 'SEP', 10: 'OCT', 11: 'NOV', 12: 'DIC'
}

def get_trimestre(mes):
    if mes in [1, 2, 3]: return "PRIMERO"
    if mes in [4, 5, 6]: return "SEGUNDO"
    if mes in [7, 8, 9]: return "TERCERO"
    return "CUARTO"

def clean_phone(phone_raw):
    """Limpia y estandariza el número telefónico a formato celular ecuatoriano (ej. 0994182398 o 983546035)."""
    if not phone_raw:
        return ""
    digits = re.sub(r'\D', '', str(phone_raw))
    if digits.startswith('593') and len(digits) >= 12:
        digits = digits[3:]
    if len(digits) > 10:
        m = re.search(r'09\d{8}', digits)
        if m:
            return m.group(0)
        digits = digits[:10]
    return digits

def extract_speed_kbps(servicio, velocidad_mbps):
    """Convierte la velocidad a Kbps."""
    if velocidad_mbps and float(velocidad_mbps) > 0:
        return int(float(velocidad_mbps) * 1000)
    if servicio:
        m_gb = re.search(r'(\d+(?:\.\d+)?)\s*(?:gbps|giga)', servicio, re.IGNORECASE)
        if m_gb:
            return int(float(m_gb.group(1)) * 1000000)
        m_mb = re.search(r'(\d+(?:\.\d+)?)\s*(?:megas|mbps|mb)', servicio, re.IGNORECASE)
        if m_mb:
            return int(float(m_mb.group(1)) * 1000)
    return 200000

def parse_datetime_flexible(val):
    """
    Parsea de forma robusta cualquier representación de fecha/hora
    (objeto datetime, date, string RFC 1123/GMT, string ISO, formato con barras, etc.)
    retornando un objeto datetime naive para Excel.
    """
    if not val:
        return None
    if isinstance(val, datetime):
        return val.replace(tzinfo=None)
    if isinstance(val, date):
        return datetime.combine(val, datetime.min.time())

    val_str = str(val).strip()
    if not val_str or val_str.lower() in ['none', 'null']:
        return None

    formatos = [
        "%a, %d %b %Y %H:%M:%S %Z",  # "Sat, 01 Aug 2026 09:20:22 GMT"
        "%a, %d %b %Y %H:%M:%S",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%dT%H:%M:%S",
        "%Y-%m-%dT%H:%M:%S.%f",
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M",
        "%d/%m/%y %H:%M:%S",
        "%d/%m/%y %H:%M",
        "%m/%d/%Y %H:%M:%S",
        "%m/%d/%Y %H:%M",
        "%Y-%m-%d"
    ]
    for fmt in formatos:
        try:
            return datetime.strptime(val_str, fmt)
        except (ValueError, TypeError):
            continue

    try:
        from dateutil import parser
        dt = parser.parse(val_str)
        return dt.replace(tzinfo=None)
    except:
        pass

    return None

# =========================================================================
# HOMOLOGACIÓN A CATÁLOGOS OFICIALES ARCOTEL
# =========================================================================

def homologar_gpon_averia(problema):
    p = (problema or '').upper().strip()
    # Descartar explícitamente cambios de fibra, configuraciones, comerciales y adicionales
    if 'CAMBIO DE FO' in p or 'CAMBIO DE FIBRA' in p or 'TENDIDO' in p:
        return None
    if 'CONF' in p or 'CONFIGURACI' in p or 'CONFIG' in p:
        return None
    if 'RETENCION' in p or 'FIDELIZACION' in p or 'COLOCAR ROUTER' in p or 'CAMBIO DE ROUTER' in p:
        return None
    if 'REUBICACION' in p or 'DOMOTICA' in p or 'COBERTURA' in p or 'INSTALACION' in p or 'VERIFICAR' in p:
        return None

    # Catálogo oficial ARCOTEL (GPON/RADIAL)
    if 'ALARMADO' in p or 'LOS' in p:
        return 'EQUIPO ALARMADO'
    if 'POTENCIA' in p:
        return 'POTENCIA DEGRADADA'
    if 'REVISION DE ONT' in p or 'CAMBIO DE ONT' in p or 'REVISION/CAMBIO DE ONT' in p or p == 'ONT':
        return 'REVISION/CAMBIO DE ONT'
    if 'INTERMITENCIA' in p or 'INTERMITENCIAS' in p:
        return 'INTERMITENCIAS EN EL SERVICIO'
    if 'ROUTER NO DA' in p or 'NO DA INTERNET' in p or 'SIN NAVEGACION' in p:
        return 'ROUTER NO DA INTERNET'
    return None

def homologar_gpon_solucion(solucion):
    """
    Soluciones oficiales GPON/RADIAL según catálogo ARCOTEL:
    1. CAMBIO DE CONECTORES APC/UPC:
       - Incluye visitas con CAMBIO DE CONECTORES APC/UPC
       - Incluye visitas con COLOCACIÓN DE UNIÓN
       - Incluye visitas con REVISIÓN EN/DE DAÑO DE FIBRA + CAMBIO DE CONECTORES
    2. CAMBIO DE CONECTORES RJ45
    3. CAMBIO DE EQUIPO ONT

    Exclusiones estrictas:
    - NO entra cambio de fibra completo (CAMBIO DE FO, tendido, etc.)
    - NO entran configuraciones de equipos (CONF. DE EQUIPOS GPON/ROUTER, etc.)
    - NO entran retenciones, fidelizaciones, routers comerciales, inspecciones, etc.
    """
    s = (solucion or '').upper().strip()

    # Descartar explícitamente cambio de fibra completo (a menos que sea revisión de daño + cambio conectores)
    if 'CAMBIO DE FO' in s or 'CAMBIO DE FIBRA' in s or 'TENDIDO' in s:
        if not (('DAÑO' in s or 'DANO' in s) and 'CONECTOR' in s):
            return None

    if 'CONF' in s or 'CONFIGURACI' in s or 'CONFIG' in s:
        return None
    if 'ARREGLO DE INSTALACI' in s or 'INSTALACI' in s or 'ADICIONAL' in s or 'INSPECCI' in s or 'REUBICACI' in s:
        return None
    if 'CAMBIO DE ROUTER' in s or 'CAMBIO DE ROUTER ANTIGUO' in s:
        return None
    if 'TICKET AL NOC' in s or 'SOLUCIÓN PARCIAL' in s or 'SOLUCION PARCIAL' in s:
        return None

    # 1. Revisión en/de daño de fibra + cambio de conectores -> CAMBIO DE CONECTORES APC/UPC
    if ('DAÑO' in s or 'DANO' in s) and 'CONECTOR' in s:
        return 'CAMBIO DE CONECTORES APC/UPC'

    # 2. Colocación de unión -> CAMBIO DE CONECTORES APC/UPC (según instrucción del usuario para el archivo de excel)
    if 'UNI' in s:
        return 'CAMBIO DE CONECTORES APC/UPC'

    # 3. Conectores APC/UPC / Conector roto -> CAMBIO DE CONECTORES APC/UPC
    if 'APC' in s or 'UPC' in s or 'CONECTOR ROTO' in s or 'CONECTORES' in s or 'CONECTOR' in s:
        if 'RJ45' in s and 'APC' not in s and 'UPC' not in s:
            return 'CAMBIO DE CONECTORES RJ45'
        return 'CAMBIO DE CONECTORES APC/UPC'

    # 4. Conectores RJ45
    if 'RJ45' in s or 'UTP' in s or 'CABLE RED' in s:
        return 'CAMBIO DE CONECTORES RJ45'

    # 5. ONT
    if 'ONT' in s or 'CAMBIO DE EQUIPO' in s:
        return 'CAMBIO DE EQUIPO ONT'

    return None

def homologar_hfc_averia(problema):
    p = (problema or '').upper().strip()
    if 'SIN SERVICIO DE CABLE' in p:
        return 'SIN SERVICIO DE CABLE'
    if 'CANALES BORROSOS' in p or 'BORROSO' in p:
        return 'CANALES BORROSOS'
    return None

def homologar_hfc_solucion(solucion):
    return 'CAMBIO DE CONECTORES RG6'

def es_paquete_hfc_valido(servicio, producto, dir_producto):
    """
    Filtro de paquete para CABLE / HFC:
    - NO van los paquetes que digan que son CABLE GPON (ni COMBO GPON, ni planes GPON de TV/Cable).
    - Si el paquete dice solo CABLE (o planes coaxiales HFC como CABLE 01, CABLE 10, CABLE 21, etc.),
      SÍ ingresa porque es HFC.
    """
    s = (servicio or '').upper().strip()
    p = (producto or '').upper().strip()
    dp = (dir_producto or '').upper().strip()

    # 1. Si el servicio o producto de la visita dice explícitamente GPON (ej. CABLE_GPON, COMBO_GPON, INTERNET_GPON)
    if 'GPON' in s or 'GPON' in p:
        return False

    # 2. Si el producto registrado en directorio_clientes menciona CABLE GPON, TV GPON o COMBO GPON
    if 'GPON' in dp:
        if any(k in dp for k in ['CABLE', 'TV', 'COMBO', 'ZAPPING']):
            return False

    # 3. Debe corresponder a servicio CABLE / HFC / COAXIAL
    texto = f"{s} {p} {dp}"
    if any(k in texto for k in ['CABLE', 'HFC', 'COAXIAL']):
        return True

    return False

def homologar_velocidad_averia(problema):
    p = (problema or '').upper().strip()
    if 'NO MARCA VELOCIDAD' in p:
        return 'NO MARCA VELOCIDAD CONTRATADA'
    return None

def homologar_velocidad_solucion(solucion):
    """
    Soluciones exactas válidas para el apartado VELOCIDAD:
    - CAMBIO DE CONECTORES RJ45
    - CAMBIO DE ONT
    Nada más.
    """
    s = (solucion or '').upper().strip()
    if 'RJ45' in s or 'UTP' in s:
        return 'CAMBIO DE CONECTORES RJ45'
    if 'ONT' in s or 'CAMBIO DE EQUIPO ONT' in s:
        return 'CAMBIO DE ONT'
    return None


# =========================================================================
# CONSULTA Y PROCESAMIENTO DE DATOS
# =========================================================================

def obtener_datos_arcotel(mes, anio, excluir_mayores_24h=True):
    """
    Obtiene las visitas del mes/año especificado y las clasifica en:
    - gpon
    - hfc
    - velocidad
    """
    conexion = db_config.get_db_connection()
    if not conexion:
        raise Exception("Error de conexión a la base de datos.")
    
    try:
        cursor = conexion.cursor(dictionary=True)
        query = """
            SELECT 
                v.id_visita, 
                v.contrato,
                v.cliente, 
                v.telefonos, 
                v.servicio, 
                v.producto,
                d.producto as dir_producto,
                v.velocidad_mbps, 
                v.problema, 
                v.solucion_tecnico, 
                v.fecha_registro, 
                v.hora_fin_visita,
                TIMESTAMPDIFF(SECOND, v.fecha_registro, v.hora_fin_visita) / 3600.0 as dur_horas
            FROM visitas_tecnicas v
            LEFT JOIN directorio_clientes d ON v.contrato = d.contrato
            WHERE v.estado = 'FINALIZADA'
              AND (v.es_instalacion = 0 OR v.es_instalacion IS NULL)
              AND YEAR(v.fecha_programada) = %s
              AND MONTH(v.fecha_programada) = %s
              AND v.hora_fin_visita IS NOT NULL
            ORDER BY v.fecha_registro ASC, v.hora_fin_visita ASC
        """
        cursor.execute(query, (int(anio), int(mes)))
        visitas = cursor.fetchall()
        cursor.close()
        conexion.close()
    except Exception as e:
        if conexion:
            conexion.close()
        raise e

    gpon_rows = []
    hfc_rows = []
    vel_rows = []

    item_gpon = 1
    item_hfc = 1
    item_vel = 1

    for v in visitas:
        dur = v['dur_horas']
        if dur is None or dur <= 0:
            continue
        
        # Filtro de 24 horas si está activo
        es_mayor_24h = (dur > 24.0)
        if excluir_mayores_24h and es_mayor_24h:
            continue

        prob = (v['problema'] or '').upper()
        serv = (v['servicio'] or '').upper()
        tel = clean_phone(v['telefonos'])
        cliente_clean = (v['cliente'] or '').strip().upper()
        f_rep_dt = parse_datetime_flexible(v['fecha_registro'])
        f_sol_dt = parse_datetime_flexible(v['hora_fin_visita'])
        f_rep = f_rep_dt.strftime('%Y-%m-%d %H:%M:%S') if f_rep_dt else ''
        f_sol = f_sol_dt.strftime('%Y-%m-%d %H:%M:%S') if f_sol_dt else ''

        # 1. VELOCIDAD (Solo problema 'NO MARCA VELOCIDAD CONTRATADA' con soluciones 'CAMBIO DE CONECTORES RJ45' o 'CAMBIO DE ONT')
        ave_vel = homologar_velocidad_averia(v['problema'])
        if ave_vel:
            sol_vel = homologar_velocidad_solucion(v['solucion_tecnico'])
            if sol_vel:
                cap_kbps = extract_speed_kbps(v['servicio'], v['velocidad_mbps'])
                obj_reclamo = int(round(cap_kbps * 0.2)) # 20% por defecto
                vel_rows.append({
                    'item': item_vel,
                    'id_visita': v['id_visita'],
                    'provincia': 'AZUAY',
                    'fecha_registro': f_rep,
                    'cliente': cliente_clean,
                    'telefonos': tel,
                    'canal': 'TELEFONICO',
                    'capacidad_kbps': cap_kbps,
                    'comparticion': '2.1',
                    'objeto_reclamo_kbps': obj_reclamo,
                    'solucion': sol_vel,
                    'solucion_original': v['solucion_tecnico'],
                    'problema_original': v['problema'],
                    'dur_horas': round(dur, 2),
                    'es_mayor_24h': es_mayor_24h
                })
                item_vel += 1
            # Si era reclamo de velocidad pero la solución no fue ni RJ45 ni ONT, se descarta
            continue

        # Si el problema era lentitud o velocidad pero no entró a VELOCIDAD, tampoco va a GPON ni HFC
        if 'VELOCIDAD' in prob or 'LENTITUD' in prob:
            continue

        # 2. CABLE / HFC (Formulario SAV-Q-001)
        # Problemas admitidos: 'SIN SERVICIO DE CABLE' y 'CANALES BORROSOS'
        # Filtro de paquete: Solo si el paquete es CABLE puro HFC. Si dice 'CABLE GPON' o tiene 'GPON', NO ENTRA.
        # Solución oficial: 'CAMBIO DE CONECTORES RG6'
        ave_hfc = homologar_hfc_averia(v['problema'])
        if ave_hfc:
            if es_paquete_hfc_valido(v.get('servicio'), v.get('producto'), v.get('dir_producto')):
                hfc_rows.append({
                    'item': item_hfc,
                    'id_visita': v['id_visita'],
                    'cliente': cliente_clean,
                    'telefonos': tel,
                    'canal': 'Via telefónica',
                    'categoria': ave_hfc,
                    'solucion': 'CAMBIO DE CONECTORES RG6',
                    'solucion_original': v['solucion_tecnico'],
                    'problema_original': v['problema'],
                    'fecha_registro': f_rep,
                    'hora_fin_visita': f_sol,
                    'dur_horas': round(dur, 2),
                    'es_mayor_24h': es_mayor_24h
                })
                item_hfc += 1
            # Si era avería de cable pero de cliente GPON, se descarta (no va ni a HFC ni a GPON)
            continue

        # 3. GPON
        else:
            ave = homologar_gpon_averia(v['problema'])
            sol = homologar_gpon_solucion(v['solucion_tecnico'])
            # Si no corresponde a una avería o solución oficial de GPON, se descarta
            if not ave or not sol:
                continue

            gpon_rows.append({
                'item': item_gpon,
                'id_visita': v['id_visita'],
                'provincia': 'AZUAY',
                'cliente': cliente_clean,
                'telefonos': tel,
                'tipo_conexion': 'NO CONMUTADA',
                'canal': 'TELEFONICO',
                'averia': ave,
                'fecha_registro': f_rep,
                'hora_fin_visita': f_sol,
                'solucion': sol,
                'solucion_original': v['solucion_tecnico'],
                'problema_original': v['problema'],
                'dur_horas': round(dur, 2),
                'es_mayor_24h': es_mayor_24h
            })
            item_gpon += 1

    return {
        'mes': int(mes),
        'anio': int(anio),
        'mes_nombre': MESES_NOMBRES.get(int(mes), ''),
        'mes_abrev': MESES_ABREV.get(int(mes), ''),
        'trimestre': get_trimestre(int(mes)),
        'excluir_mayores_24h': excluir_mayores_24h,
        'totales': {
            'gpon': len(gpon_rows),
            'hfc': len(hfc_rows),
            'velocidad': len(vel_rows),
            'total_general': len(gpon_rows) + len(hfc_rows) + len(vel_rows)
        },
        'gpon': gpon_rows,
        'hfc': hfc_rows,
        'velocidad': vel_rows
    }


# =========================================================================
# GENERADORES DE ARCHIVOS EXCEL
# =========================================================================

BORDER_THIN = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

BORDER_HEADER = Border(
    left=Side(style='thin', color='000000'),
    right=Side(style='thin', color='000000'),
    top=Side(style='thin', color='000000'),
    bottom=Side(style='thin', color='000000')
)

def generar_excel_arcotel_gpon(filas, mes, anio):
    """Genera el Excel oficial de GPON idéntico a GPON-JULIO-2026.xlsx."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    ws.views.sheetView[0].showGridLines = True

    # Fila 1: Título Principal
    ws.merge_cells('A1:K1')
    cell_a1 = ws['A1']
    cell_a1.value = "TIEMPO PROMEDIO DE REPARACIÓN DE AVERÍAS (TÉCNICAS)"
    cell_a1.font = Font(name='Calibri', size=11, bold=True)
    cell_a1.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 24

    # Fila 2: Sub-bloques
    ws.merge_cells('A2:E2')
    ws['A2'].value = "1) DATOS DEL INGRESO DEL RECLAMO"
    ws['A2'].font = Font(name='Calibri', size=10, bold=True)
    ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
    ws['A2'].fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')

    ws.merge_cells('F2:K2')
    ws['F2'].value = "2) DETALLES DEL RECLAMO"
    ws['F2'].font = Font(name='Calibri', size=10, bold=True)
    ws['F2'].alignment = Alignment(horizontal='center', vertical='center')
    ws['F2'].fill = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')
    ws.row_dimensions[2].height = 20

    for col in range(1, 12):
        cell = ws.cell(2, col)
        cell.border = BORDER_HEADER

    # Fila 3: Encabezados Oficiales
    headers = [
        "ITEM",
        "PROVINCIA",
        "NOMBRE DE LA PERSONA QUE REALIZA EL REQUERIMIENTO",
        "NÚMERO TELEFÓNICO DE CONTACTO DEL USUARIO",
        "TIPO DE CONEXIÓN (CONMUTADA O NO CONMUTADA)",
        "CANAL DE REQUERIMIENTO (PERSONALIZADO, TELEFÓNICO, OFICIO, CORREO ELECTRÓNICO, PÁGINA WEB)",
        "TIPO DE AVERÍA",
        "FECHA Y HORA DE REPORTE DE LA AVERÍA (dd/mm/aaaa hh:mm)",
        "FECHA Y HORA DE REPARACIÓN DE LA AVERÍA (dd/mm/aaaa hh:mm)",
        "TIEMPO DE REPARACIÓN DE LA AVERÍA (calculo en HORAS) ( Campo No obligatorio)",
        "DESCRIPCIÓN DE LA SOLUCIÓN"
    ]

    ws.row_dimensions[3].height = 45
    for c_idx, h in enumerate(headers, 1):
        cell = ws.cell(3, c_idx)
        cell.value = h
        cell.font = Font(name='Calibri', size=9, bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER_HEADER
        if c_idx <= 5:
            cell.fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
        else:
            cell.fill = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')

    # Filas de datos
    for r_idx, f in enumerate(filas, 4):
        ws.row_dimensions[r_idx].height = 20
        sol_val = f.get('solucion', '')
        sol_u = (sol_val or '').upper()
        if 'UNI' in sol_u or (('DAÑO' in sol_u or 'DANO' in sol_u) and 'CONECTOR' in sol_u):
            sol_val = 'CAMBIO DE CONECTORES APC/UPC'

        f_rep = parse_datetime_flexible(f.get('fecha_registro'))
        f_sol = parse_datetime_flexible(f.get('hora_fin_visita'))
        formula_dur = f"=I{r_idx}-H{r_idx}" if (f_rep and f_sol) else ""

        row_vals = [
            f.get('item', r_idx - 3),
            f.get('provincia', 'AZUAY'),
            f.get('cliente', ''),
            f.get('telefonos', ''),
            f.get('tipo_conexion', 'NO CONMUTADA'),
            f.get('canal', 'TELEFONICO'),
            f.get('averia', ''),
            f_rep,
            f_sol,
            formula_dur,
            sol_val
        ]

        for c_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(r_idx, c_idx)
            cell.value = val
            cell.font = Font(name='Calibri', size=10)
            cell.border = BORDER_THIN

            if c_idx in [1, 2, 5, 6, 7, 11]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif c_idx == 3:
                cell.alignment = Alignment(horizontal='left', vertical='center')
            elif c_idx == 4:
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.number_format = '0000000000'
            elif c_idx in [8, 9]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.number_format = 'd/m/yyyy h:mm'
            elif c_idx == 10:
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.number_format = '[h]:mm:ss'

    # Anchos de columna
    col_widths = {
        'A': 6.0, 'B': 12.0, 'C': 42.0, 'D': 15.0, 'E': 18.0,
        'F': 20.0, 'G': 32.0, 'H': 23.0, 'I': 23.0, 'J': 24.0, 'K': 45.0
    }
    for col_letter, w in col_widths.items():
        ws.column_dimensions[col_letter].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def generar_excel_arcotel_hfc(filas, mes, anio):
    """Genera el Excel oficial de HFC idéntico a HFC_MAY_2026.xlsx."""
    wb = openpyxl.Workbook()
    ws = wb.active
    mes_str = MESES_NOMBRES.get(int(mes), 'MES')
    ws.title = f"HFC {mes_str} {anio}"
    ws.views.sheetView[0].showGridLines = True

    # Fila 1 y 2: Título institucional
    ws.merge_cells('D1:H1')
    ws['D1'].value = "FORMULARIO PARA EL ANÁLISIS DE LA CALIDAD DE SERVICIO: "
    ws['D1'].font = Font(name='Calibri', size=12, bold=True)
    ws['D1'].alignment = Alignment(horizontal='center', vertical='center')

    ws.merge_cells('D2:H2')
    ws['D2'].value = "DETALLE DE REQUERIMIENTOS DE ATENCIÓN DE LOS SUSCRIPTORES"
    ws['D2'].font = Font(name='Calibri', size=12, bold=True)
    ws['D2'].alignment = Alignment(horizontal='center', vertical='center')

    # Fila 4: Código Formulario
    ws['B4'].value = "SAV-Q-001"
    ws['B4'].font = Font(name='Helvetica Neue', size=10, bold=True)

    # Fila 5: Año
    ws['A5'].value = f"AÑO: {anio}"
    ws['A5'].font = Font(name='Helvetica Neue', size=10, bold=True)

    # Fila 6: Trimestre
    trimestre = get_trimestre(int(mes))
    ws['A6'].value = f"TRIMESTRE: {trimestre}"
    ws['A6'].font = Font(name='Helvetica Neue', size=10, bold=True)

    # Fila 7: Encabezados (8 columnas)
    headers = [
        "N°. Item",
        "Nombre y Apellido del Suscriptor",
        "Teléfono",
        "Forma de reclamo (Vía telefónica, Correo electrónico, Oficio o carta, Ventanilla de agencia o caja, Centro de atención ARCOTEL )",
        "Categoría (Avería técnica, Servicio al cliente, Reclamo de facturación)",
        "Descripción del requerimiento de atención",
        "Fecha y Hora de ingreso del requerimiento de atención (dd/mm/aaaa hh:mm)",
        "Fecha y Hora de solución del requerimiento de atención (dd/mm/aaaa hh:mm)"
    ]

    ws.row_dimensions[7].height = 45
    for c_idx, h in enumerate(headers, 1):
        cell = ws.cell(7, c_idx)
        cell.value = h
        cell.font = Font(name='Helvetica Neue', size=9, bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER_HEADER
        cell.fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')

    # Filas de datos
    for r_idx, f in enumerate(filas, 8):
        ws.row_dimensions[r_idx].height = 20
        tel_val = str(f.get('telefonos') or '')
        if tel_val and not tel_val.startswith('0') and len(tel_val) == 9:
            tel_val = '0' + tel_val

        f_ing = parse_datetime_flexible(f.get('fecha_registro'))
        f_sol = parse_datetime_flexible(f.get('hora_fin_visita'))

        row_vals = [
            f.get('item', r_idx - 7),
            f.get('cliente', ''),
            tel_val,
            f.get('canal', 'Via telefónica'),
            f.get('categoria', 'SIN SERVICIO DE CABLE'),
            f.get('solucion', 'CAMBIO DE CONECTORES RG6'),
            f_ing,
            f_sol
        ]

        for c_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(r_idx, c_idx)
            cell.value = val
            cell.font = Font(name='Calibri', size=10)
            cell.border = BORDER_THIN

            if c_idx in [1, 4, 5, 6]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif c_idx == 2:
                cell.alignment = Alignment(horizontal='left', vertical='center')
            elif c_idx == 3:
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.number_format = '@'
            elif c_idx in [7, 8]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.number_format = 'd/m/yyyy h:mm'

    col_widths = {
        'A': 8.0, 'B': 40.0, 'C': 15.0, 'D': 24.0,
        'E': 26.0, 'F': 32.0, 'G': 25.0, 'H': 25.0
    }
    for col_letter, w in col_widths.items():
        ws.column_dimensions[col_letter].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def generar_excel_arcotel_velocidad(filas, mes, anio):
    """Genera el Excel oficial de VELOCIDAD idéntico a VELOCIDAD-ABRIL-2026.xlsx."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    ws.views.sheetView[0].showGridLines = True

    # Fila 1: Título Principal
    ws.merge_cells('A1:J1')
    cell_a1 = ws['A1']
    cell_a1.value = "PORCENTAJE DE RECLAMOS POR LA CAPACIDAD DEL CANAL DE ACCESO CONTRATADO POR EL CLIENTE"
    cell_a1.font = Font(name='Calibri', size=11, bold=True)
    cell_a1.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 24

    # Fila 2: Sub-bloques
    ws.merge_cells('A2:E2')
    ws['A2'].value = "1) DATOS DEL INGRESO DEL RECLAMO"
    ws['A2'].font = Font(name='Calibri', size=9, bold=True)
    ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
    ws['A2'].fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')

    ws.merge_cells('F2:J2')
    ws['F2'].value = "2) DETALLES DEL RECLAMO"
    ws['F2'].font = Font(name='Calibri', size=9, bold=True)
    ws['F2'].alignment = Alignment(horizontal='center', vertical='center')
    ws['F2'].fill = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')
    ws.row_dimensions[2].height = 20

    for col in range(1, 11):
        cell = ws.cell(2, col)
        cell.border = BORDER_HEADER

    # Fila 3: Encabezados (10 columnas)
    headers = [
        "ITEM",
        "PROVINCIA",
        "FECHA Y HORA DEL REGISTRO DEL RECLAMO",
        "NOMBRE DE LA PERSONA QUE REALIZA EL RECLAMO",
        "NÚMERO TELEFÓNICO DE CONTACTO DEL USUARIO",
        "CANAL DE RECLAMO (PERSONALIZADO, TELEFÓNICO, CORREO ELECTRÓNICO, PÁGINA WEB U OFICIO)",
        "CAPACIDAD EFECTIVA CONTRATADA (Kbps)",
        "COMPARTICIÓN",
        "% DE CAPACIDAD EFECTIVAMENTE SUMINISTRADO (OBJETO DEL RECLAMO)",
        "DESCRIPCIÓN DE LA SOLUCIÓN"
    ]

    ws.row_dimensions[3].height = 45
    for c_idx, h in enumerate(headers, 1):
        cell = ws.cell(3, c_idx)
        cell.value = h
        cell.font = Font(name='Calibri', size=9, bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER_HEADER
        if c_idx <= 5:
            cell.fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
        else:
            cell.fill = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')

    # Filas de datos
    for r_idx, f in enumerate(filas, 4):
        ws.row_dimensions[r_idx].height = 20
        tel_val = str(f.get('telefonos') or '')
        if tel_val and not tel_val.startswith('0') and len(tel_val) == 9:
            tel_val = '0' + tel_val

        cap_kbps = f.get('capacidad_kbps', 200000)
        obj_reclamo = f.get('objeto_reclamo_kbps', int(round(cap_kbps * 0.2)))

        f_reg = parse_datetime_flexible(f.get('fecha_registro'))

        row_vals = [
            f.get('item', r_idx - 3),
            f.get('provincia', 'AZUAY'),
            f_reg,
            f.get('cliente', ''),
            tel_val,
            f.get('canal', 'TELEFONICO'),
            cap_kbps,
            f.get('comparticion', '2.1'),
            obj_reclamo,
            f.get('solucion', 'CAMBIO DE CONECTORES RJ45')
        ]

        for c_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(r_idx, c_idx)
            cell.value = val
            cell.font = Font(name='Calibri', size=10)
            cell.border = BORDER_THIN

            if c_idx in [1, 2, 6, 8, 10]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif c_idx == 3:
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.number_format = 'd/m/yyyy h:mm'
            elif c_idx == 4:
                cell.alignment = Alignment(horizontal='left', vertical='center')
            elif c_idx == 5:
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.number_format = '@'
            elif c_idx in [7, 9]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.number_format = '#,##0'

    col_widths = {
        'A': 7.0, 'B': 14.0, 'C': 24.0, 'D': 38.0, 'E': 16.0,
        'F': 22.0, 'G': 20.0, 'H': 15.0, 'I': 22.0, 'J': 30.0
    }
    for col_letter, w in col_widths.items():
        ws.column_dimensions[col_letter].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def generar_zip_arcotel(mes, anio, excluir_mayores_24h=True):
    """Genera un archivo ZIP con los 3 reportes oficiales de ARCOTEL."""
    data = obtener_datos_arcotel(mes, anio, excluir_mayores_24h)
    mes_str = MESES_NOMBRES.get(int(mes), 'MES')
    mes_abrev = MESES_ABREV.get(int(mes), 'MES')

    buf_gpon = generar_excel_arcotel_gpon(data['gpon'], mes, anio)
    buf_hfc = generar_excel_arcotel_hfc(data['hfc'], mes, anio)
    buf_vel = generar_excel_arcotel_velocidad(data['velocidad'], mes, anio)

    zip_buf = io.BytesIO()
    with zipfile.ZipFile(zip_buf, mode='w', compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(f"GPON-{mes_str}-{anio}.xlsx", buf_gpon.getvalue())
        zf.writestr(f"HFC_{mes_abrev}_{anio}.xlsx", buf_hfc.getvalue())
        zf.writestr(f"VELOCIDAD-{mes_str}-{anio}.xlsx", buf_vel.getvalue())

    zip_buf.seek(0)
    return zip_buf
